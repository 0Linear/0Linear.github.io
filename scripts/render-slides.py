#!/usr/bin/env python3
"""
Turns a PowerPoint deck into the images the presentation viewer shows.

    python3 scripts/render-slides.py path/to/deck.pptx

Writes presentation/slides/ (full-size WebP, 2560px wide), presentation/thumbs/
(small WebP) and presentation/slides.json (slide titles and sizes). Commit those;
the site build only copies them, so it needs no LibreOffice.

Needs: LibreOffice (soffice), Poppler (pdftoppm), Python packages Pillow and
python-pptx. Install the fonts the deck uses first, or LibreOffice substitutes
others. Set SOFFICE to use a different LibreOffice command.
"""
import glob
import json
import os
import shlex
import shutil
import subprocess
import sys
import re
import tempfile
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'presentation')
WIDTH = 2560
THUMB_WIDTH = 480


def slide_title(slide):
    """The title placeholder's text, else the text set in the largest font."""
    if slide.shapes.title is not None and slide.shapes.title.text.strip():
        return ' '.join(slide.shapes.title.text.split())

    best, best_size = '', 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            text = ' '.join(''.join(r.text for r in para.runs).split())
            sizes = [r.font.size for r in para.runs if r.font.size]
            size = max(sizes) if sizes else Pt(18)
            if text and size > best_size:
                best, best_size = text, size
    return best


def libreoffice_friendly_copy(deck, dest):
    """
    LibreOffice centers the text of left-aligned "do not wrap" text boxes on
    the box, where PowerPoint starts it at the box's left edge. Unfilled boxes
    like that are turned into wrapping boxes reaching the right edge of the
    slide, which PowerPoint would draw the same way.
    """
    src = zipfile.ZipFile(deck)
    slide_w = int(re.search(r'<p:sldSz cx="(\d+)"', src.read('ppt/presentation.xml').decode()).group(1))

    def fix_shape(m):
        sp = m.group(0)
        if 'wrap="none"' not in sp or re.search(r'<a:(solidFill|gradFill|blipFill)', sp.split('<p:txBody>')[0]):
            return sp
        if re.search(r'algn="(ctr|r|just|dist)"', sp):
            return sp
        off = re.search(r'<a:off x="(-?\d+)"', sp)
        ext = re.search(r'<a:ext cx="(\d+)"', sp)
        if not off or not ext:
            return sp
        width = max(int(ext.group(1)), slide_w - int(off.group(1)))
        sp = sp.replace('wrap="none"', 'wrap="square"', 1)
        return sp[:ext.start(1)] + str(width) + sp[ext.end(1):]

    with zipfile.ZipFile(dest, 'w', zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if re.match(r'ppt/slides/slide\d+\.xml$', item.filename):
                data = re.sub(r'<p:sp>.*?</p:sp>', fix_shape, data.decode('utf-8'), flags=re.S).encode('utf-8')
            out.writestr(item, data)


def main(deck):
    prs = Presentation(deck)
    titles = [slide_title(s) for s in prs.slides]
    deck_title = prs.core_properties.title or (titles[0] if titles else 'Presentation')
    inches_wide = prs.slide_width / 914400

    soffice = shlex.split(os.environ.get('SOFFICE', 'soffice'))
    with tempfile.TemporaryDirectory() as tmp:
        fixed = os.path.join(tmp, 'deck.pptx')
        libreoffice_friendly_copy(deck, fixed)
        subprocess.run(soffice + ['--headless', '--convert-to', 'pdf', '--outdir', tmp, fixed], check=True)
        pdf = glob.glob(os.path.join(tmp, '*.pdf'))[0]
        dpi = round(WIDTH / inches_wide) + 1
        height = round(WIDTH * prs.slide_height / prs.slide_width)
        subprocess.run(['pdftoppm', '-png', '-r', str(dpi), pdf, os.path.join(tmp, 'p')], check=True)
        pages = sorted(glob.glob(os.path.join(tmp, 'p-*.png')),
                       key=lambda p: int(p.rsplit('-', 1)[1].split('.')[0]))

        for d in ('slides', 'thumbs'):
            shutil.rmtree(os.path.join(OUT, d), ignore_errors=True)
            os.makedirs(os.path.join(OUT, d))

        slides = []
        for i, page in enumerate(pages, 1):
            name = f'slide-{i:02d}.webp'
            img = Image.open(page).convert('RGB')
            # pdftoppm rounds the page size and can leave a blank last row or
            # column; trim one pixel and scale to the slide's exact shape
            img = img.crop((0, 0, img.width - 1, img.height - 1)).resize((WIDTH, height), Image.LANCZOS)
            img.save(os.path.join(OUT, 'slides', name), 'WEBP', quality=90, method=6)
            thumb = img.resize((THUMB_WIDTH, round(img.height * THUMB_WIDTH / img.width)), Image.LANCZOS)
            thumb.save(os.path.join(OUT, 'thumbs', name), 'WEBP', quality=82, method=6)
            slides.append({
                'src': f'slides/{name}',
                'thumb': f'thumbs/{name}',
                'title': titles[i - 1] if i <= len(titles) else f'Slide {i}',
                'width': img.width,
                'height': img.height
            })

    with open(os.path.join(OUT, 'slides.json'), 'w') as f:
        json.dump({'title': deck_title, 'slides': slides}, f, indent=2, ensure_ascii=False)
        f.write('\n')
    print(f'{len(slides)} slides -> {os.path.relpath(OUT)}')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(os.path.abspath(sys.argv[1]))
